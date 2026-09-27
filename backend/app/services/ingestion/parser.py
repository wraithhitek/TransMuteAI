import io
import uuid
from typing import Tuple, List
from pathlib import Path
from pypdf import PdfReader
from docx import Document

from app.services.ingestion.pii_redactor import PIIRedactor
from app.services.ingestion.ocr_extractor import OCRExtractor
from app.api.schemas import IngestionResponse, RedactedItem

class DocumentParser:
    """
    Multi-modal parser converting PDF, DOCX, TXT, and Images into clean, redacted text.
    """

    SUPPORTED_EXTENSIONS = {
        ".pdf": "PDF",
        ".docx": "DOCX",
        ".txt": "TEXT",
        ".md": "MARKDOWN",
        ".png": "IMAGE",
        ".jpg": "IMAGE",
        ".jpeg": "IMAGE",
        ".webp": "IMAGE",
        ".tiff": "IMAGE",
    }

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> str:
        reader = PdfReader(io.BytesIO(file_bytes))
        extracted_pages = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                extracted_pages.append(text.strip())
        return "\n\n".join(extracted_pages)

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> str:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        return "\n\n".join(paragraphs)

    @classmethod
    def parse_image(cls, file_bytes: bytes) -> str:
        return OCRExtractor.extract_text(file_bytes)

    @classmethod
    def parse_file(cls, filename: str, file_bytes: bytes) -> IngestionResponse:
        ext = Path(filename).suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file format '{ext}'. Supported: {list(cls.SUPPORTED_EXTENSIONS.keys())}")

        file_type = cls.SUPPORTED_EXTENSIONS[ext]
        raw_text = ""

        if file_type == "PDF":
            raw_text = cls.parse_pdf(file_bytes)
        elif file_type == "DOCX":
            raw_text = cls.parse_docx(file_bytes)
        elif file_type == "IMAGE":
            raw_text = cls.parse_image(file_bytes)
        else:
            # Text / Markdown
            raw_text = file_bytes.decode("utf-8", errors="replace")

        if not raw_text.strip():
            raw_text = f"[Empty document or no readable text extracted from {filename}]"

        # Apply PII Redaction
        redacted_text, redactions = PIIRedactor.redact(raw_text)

        preview = redacted_text[:500] + ("..." if len(redacted_text) > 500 else "")

        return IngestionResponse(
            document_id=f"doc_{uuid.uuid4().hex[:12]}",
            filename=filename,
            file_type=file_type,
            raw_character_count=len(raw_text),
            redacted_character_count=len(redacted_text),
            redaction_count=len(redactions),
            redactions=redactions,
            redacted_text=redacted_text,
            preview=preview
        )
