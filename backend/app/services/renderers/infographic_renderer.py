import os
from typing import Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.api.schemas import ContentBriefJSON
from app.services.renderers.base import BaseRenderer

class InfographicRenderer(BaseRenderer):
    """
    Renders Infographic content, visual layout recommendations, and key messaging:
    1. Comprehensive Infographic Design Blueprint (DOCX)
    2. Interactive HTML Visual Layout Blueprint (.html)
    """

    PRIMARY_COLOR = RGBColor(15, 23, 42)      # Slate 900
    ACCENT_COLOR = RGBColor(6, 182, 212)      # Cyan 500
    MUTED_COLOR = RGBColor(100, 116, 139)     # Slate 500

    def _set_cell_background(self, cell, hex_color: str):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def render(self, brief: ContentBriefJSON, output_path: str, extra_context: Optional[Dict[str, Any]] = None) -> str:
        info = brief.infographic
        doc = Document()

        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Header
        title_para = doc.add_paragraph()
        run_title = title_para.add_run(f"Infographic Blueprint: {info.headline if info and info.headline else brief.title}")
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(22)
        run_title.font.bold = True
        run_title.font.color.rgb = self.PRIMARY_COLOR

        sub_para = doc.add_paragraph()
        flow = info.narrative_flow if info else "Problem ➔ Solution ➔ Metrics ➔ Impact"
        run_sub = sub_para.add_run(f"Visual Layout Recommendations & Narrative Flow: {flow}")
        run_sub.font.size = Pt(11)
        run_sub.font.bold = True
        run_sub.font.color.rgb = self.ACCENT_COLOR

        # Color Palette Directive Box
        palette = info.color_palette if (info and info.color_palette) else ["#0F172A", "#2563EB", "#06B6D4", "#10B981"]
        pal_table = doc.add_table(rows=1, cols=len(palette))
        pal_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, hex_code in enumerate(palette):
            cell = pal_table.rows[0].cells[idx]
            clean_hex = hex_code.replace("#", "")
            try:
                self._set_cell_background(cell, clean_hex)
            except Exception:
                self._set_cell_background(cell, "0F172A")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"\n{hex_code}\n")
            r.font.size = Pt(8.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # Quadrant Sections Table
        h_sec = doc.add_paragraph()
        r_sec = h_sec.add_run("Infographic Content Quadrants & Visual Directives")
        r_sec.font.size = Pt(14)
        r_sec.font.bold = True
        r_sec.font.color.rgb = self.PRIMARY_COLOR

        sections = info.sections if (info and info.sections) else []
        if sections:
            table = doc.add_table(rows=1, cols=4)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["Section Title", "Hero Metric / Stat", "Visual Layout Directive", "Supporting Narrative"]
            for idx, h_text in enumerate(headers):
                cell = table.rows[0].cells[idx]
                self._set_cell_background(cell, "1E293B")
                p = cell.paragraphs[0]
                r = p.add_run(h_text)
                r.font.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(255, 255, 255)

            for s in sections:
                row = table.add_row()
                # Col 0: Section
                self._set_cell_background(row.cells[0], "F8FAFC")
                p0 = row.cells[0].paragraphs[0]
                r0 = p0.add_run(s.section_title)
                r0.font.bold = True
                r0.font.size = Pt(9.5)

                # Col 1: Key Stat
                self._set_cell_background(row.cells[1], "ECFEFF")  # Cyan-50
                p1 = row.cells[1].paragraphs[0]
                p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r1 = p1.add_run(s.key_stat)
                r1.font.bold = True
                r1.font.size = Pt(12)
                r1.font.color.rgb = RGBColor(8, 145, 178)

                # Col 2: Visual Directive
                self._set_cell_background(row.cells[2], "FFFFFF")
                p2 = row.cells[2].paragraphs[0]
                r2 = p2.add_run(s.visual_layout_directive)
                r2.font.size = Pt(9)
                r2.font.italic = True

                # Col 3: Supporting Text
                self._set_cell_background(row.cells[3], "F8FAFC")
                p3 = row.cells[3].paragraphs[0]
                r3 = p3.add_run(s.supporting_text)
                r3.font.size = Pt(9)

        # Footer attribution
        doc.add_paragraph().paragraph_format.space_after = Pt(12)
        callout = info.footer_callout if (info and info.footer_callout) else "Verified via TransmuteAI Cryptographic Provenance"
        p_ft = doc.add_paragraph(f"Graphic Footer Attribution: {callout}")
        p_ft.runs[0].font.size = Pt(9)
        p_ft.runs[0].font.italic = True
        p_ft.runs[0].font.color.rgb = self.MUTED_COLOR

        # QR Code
        qr_path = extra_context.get("qr_code_path") if extra_context else None
        if qr_path and os.path.exists(qr_path):
            doc.add_paragraph().paragraph_format.space_after = Pt(8)
            qr_para = doc.add_paragraph()
            qr_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            qr_para.add_run().add_picture(qr_path, width=Inches(1.2))
            cap = doc.add_paragraph("Scan to Verify Infographic Provenance")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].font.size = Pt(8)
            cap.runs[0].font.italic = True
            cap.runs[0].font.color.rgb = self.MUTED_COLOR

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        doc.save(output_path)

        # Also generate an interactive HTML preview file alongside
        html_preview_path = output_path.replace(".docx", "_preview.html")
        sections_html = "".join([
            f"""
            <div style="background:#1e293b; padding:20px; border-radius:12px; border:1px solid #334155; margin-bottom:16px;">
                <span style="font-size:11px; text-transform:uppercase; color:#38bdf8; font-weight:700;">{s.section_title}</span>
                <div style="font-size:28px; font-weight:800; color:#38bdf8; margin:8px 0;">{s.key_stat}</div>
                <div style="font-size:12px; color:#94a3b8; font-style:italic; margin-bottom:8px;">📐 Visual Layout: {s.visual_layout_directive}</div>
                <p style="font-size:13px; color:#e2e8f0; line-height:1.5;">{s.supporting_text}</p>
            </div>
            """ for s in sections
        ])

        html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Infographic Blueprint - {brief.title}</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; background: #0f172a; color: #f8fafc; padding: 40px; margin: 0; }}
        .container {{ max-width: 800px; margin: 0 auto; }}
        .header {{ text-align: center; margin-bottom: 32px; }}
        h1 {{ font-size: 26px; color: #38bdf8; margin-bottom: 8px; }}
        .flow {{ font-size: 13px; color: #94a3b8; font-weight: 600; text-transform: uppercase; }}
        .palette {{ display: flex; gap: 8px; justify-content: center; margin: 20px 0; }}
        .swatch {{ padding: 6px 14px; border-radius: 6px; font-size: 11px; font-weight: bold; color: white; border: 1px solid rgba(255,255,255,0.2); }}
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <span class="flow">{flow}</span>
            <h1>{info.headline if info and info.headline else brief.title}</h1>
            <div class="palette">
                {' '.join(f'<div class="swatch" style="background:{c};">{c}</div>' for c in palette)}
            </div>
        </div>
        <div class="quadrants">
            {sections_html}
        </div>
    </div>
</body>
</html>"""
        with open(html_preview_path, "w", encoding="utf-8") as f:
            f.write(html_content)

        return os.path.abspath(output_path)
