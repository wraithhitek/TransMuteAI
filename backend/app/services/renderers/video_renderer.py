import os
from typing import Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.api.schemas import ContentBriefJSON
from app.services.renderers.base import BaseRenderer

class VideoRenderer(BaseRenderer):
    """
    Renders complete Video Production Packages:
    1. Production Brief & Storyboard Document (DOCX)
    2. Broadcast SubRip Subtitle file (.srt)
    """

    PRIMARY_COLOR = RGBColor(15, 23, 42)      # Slate 900
    ACCENT_COLOR = RGBColor(147, 51, 234)     # Purple 600
    MUTED_COLOR = RGBColor(100, 116, 139)     # Slate 500

    def _set_cell_background(self, cell, hex_color: str):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def render(self, brief: ContentBriefJSON, output_path: str, extra_context: Optional[Dict[str, Any]] = None) -> str:
        video_data = brief.video_package
        doc = Document()

        # Page setup
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run(f"Video Production Package: {brief.title}")
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(22)
        run_title.font.bold = True
        run_title.font.color.rgb = self.PRIMARY_COLOR

        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(f"Creative Concept & Production Storyboard | Target Duration: {video_data.target_duration if video_data else '60-90s'}")
        run_sub.font.size = Pt(11)
        run_sub.font.bold = True
        run_sub.font.color.rgb = self.ACCENT_COLOR

        # Creative Concept Box
        concept_table = doc.add_table(rows=1, cols=1)
        concept_cell = concept_table.rows[0].cells[0]
        self._set_cell_background(concept_cell, "F3E8FF")
        p_c = concept_cell.paragraphs[0]
        p_c.paragraph_format.space_before = Pt(6)
        p_c.paragraph_format.space_after = Pt(6)
        r_c_lbl = p_c.add_run("DIRECTOR'S CONCEPT & VISUAL TREATMENT\n")
        r_c_lbl.font.size = Pt(9)
        r_c_lbl.font.bold = True
        r_c_lbl.font.color.rgb = self.ACCENT_COLOR
        r_c_val = p_c.add_run(video_data.concept if video_data else brief.summary)
        r_c_val.font.size = Pt(10.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # Storyboard Header
        h_sb = doc.add_paragraph()
        r_sb = h_sb.add_run("Scene-by-Scene Storyboard & Narration Script")
        r_sb.font.size = Pt(14)
        r_sb.font.bold = True
        r_sb.font.color.rgb = self.PRIMARY_COLOR

        scenes = video_data.storyboard if (video_data and video_data.storyboard) else []
        if scenes:
            table = doc.add_table(rows=1, cols=4)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["Scene / Time", "Visual Description & Cinematography", "Voiceover Narration", "Audio Cues & B-Roll"]
            for idx, h_text in enumerate(headers):
                cell = table.rows[0].cells[idx]
                self._set_cell_background(cell, "1E293B")
                p = cell.paragraphs[0]
                r = p.add_run(h_text)
                r.font.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(255, 255, 255)

            for scene in scenes:
                row = table.add_row()
                # Col 0: Scene & Time
                self._set_cell_background(row.cells[0], "F8FAFC")
                p0 = row.cells[0].paragraphs[0]
                r0_1 = p0.add_run(f"Scene {scene.scene_number}\n")
                r0_1.font.bold = True
                r0_1.font.size = Pt(9.5)
                r0_2 = p0.add_run(scene.timestamp)
                r0_2.font.size = Pt(8.5)
                r0_2.font.color.rgb = self.MUTED_COLOR

                # Col 1: Visuals
                self._set_cell_background(row.cells[1], "FFFFFF")
                p1 = row.cells[1].paragraphs[0]
                r1 = p1.add_run(scene.visual_description)
                r1.font.size = Pt(9.5)

                # Col 2: Narration
                self._set_cell_background(row.cells[2], "F8FAFC")
                p2 = row.cells[2].paragraphs[0]
                r2 = p2.add_run(f'"{scene.narration}"')
                r2.font.size = Pt(9.5)
                r2.font.italic = True

                # Col 3: Audio & B-Roll
                self._set_cell_background(row.cells[3], "FFFFFF")
                p3 = row.cells[3].paragraphs[0]
                r3_a = p3.add_run(f"SFX/Music: {scene.audio_cues}\n")
                r3_a.font.size = Pt(8.5)
                r3_a.font.color.rgb = self.ACCENT_COLOR
                if scene.b_roll_recommendations:
                    r3_b = p3.add_run("B-Roll: " + ", ".join(scene.b_roll_recommendations))
                    r3_b.font.size = Pt(8.5)
                    r3_b.font.color.rgb = self.MUTED_COLOR

        # Subtitle Section
        doc.add_paragraph().paragraph_format.space_after = Pt(12)
        h_sub = doc.add_paragraph()
        r_h_sub = h_sub.add_run("Broadcast Subtitles (SRT Stream)")
        r_h_sub.font.size = Pt(13)
        r_h_sub.font.bold = True
        r_h_sub.font.color.rgb = self.PRIMARY_COLOR

        srt_content = video_data.subtitles_srt if (video_data and video_data.subtitles_srt) else "1\n00:00:00,000 --> 00:00:05,000\n[Narration placeholder]\n"
        sub_para = doc.add_paragraph(srt_content)
        sub_para.paragraph_format.space_after = Pt(12)

        # Verification QR Code
        qr_path = extra_context.get("qr_code_path") if extra_context else None
        if qr_path and os.path.exists(qr_path):
            doc.add_paragraph().paragraph_format.space_after = Pt(8)
            qr_para = doc.add_paragraph()
            qr_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            qr_para.add_run().add_picture(qr_path, width=Inches(1.3))
            cap = doc.add_paragraph("Scan to Verify Video Package Provenance")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].font.size = Pt(8)
            cap.runs[0].font.italic = True
            cap.runs[0].font.color.rgb = self.MUTED_COLOR

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        doc.save(output_path)

        # Also write standalone .srt file alongside DOCX
        srt_file_path = output_path.replace(".docx", "_subtitles.srt")
        with open(srt_file_path, "w", encoding="utf-8") as f:
            f.write(srt_content)

        return os.path.abspath(output_path)
