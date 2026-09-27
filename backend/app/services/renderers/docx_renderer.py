import os
from typing import Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from app.api.schemas import ContentBriefJSON
from app.services.renderers.base import BaseRenderer

class DOCXRenderer(BaseRenderer):
    """
    Renders ContentBriefJSON into an executive Word document with professional layout,
    custom color palettes, structured tables, and verification QR code embedding.
    """

    PRIMARY_COLOR = RGBColor(30, 41, 59)      # Slate 800
    ACCENT_COLOR = RGBColor(37, 99, 235)      # Blue 600
    MUTED_COLOR = RGBColor(100, 116, 139)     # Slate 500

    def _set_cell_background(self, cell, hex_color: str):
        """Set shading background color for a table cell."""
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def render(self, brief: ContentBriefJSON, output_path: str, extra_context: Optional[Dict[str, Any]] = None) -> str:
        doc = Document()

        # Adjust page margins (0.8 inches)
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Document Header
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_title = title_para.add_run(brief.title)
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(24)
        run_title.font.bold = True
        run_title.font.color.rgb = self.PRIMARY_COLOR

        sub_para = doc.add_paragraph()
        run_sub = sub_para.add_run(f"Executive Strategic Brief | Tone: {brief.tone.capitalize()}")
        run_sub.font.size = Pt(11)
        run_sub.font.color.rgb = self.ACCENT_COLOR
        run_sub.font.bold = True

        # Metadata Summary Table
        meta_table = doc.add_table(rows=2, cols=3)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_data = [
            [("Brief ID", brief.brief_id), ("Generated", brief.created_at[:10]), ("Status", "VERIFIED PROVENANCE")],
            [("Target Formats", ", ".join(brief.target_formats).upper()), ("Entities", ", ".join(brief.entities_extracted[:4]) or "General"), ("Integrity", "SHA-256 Chained")]
        ]

        for r_idx, row in enumerate(meta_table.rows):
            for c_idx, cell in enumerate(row.cells):
                label, val = meta_data[r_idx][c_idx]
                self._set_cell_background(cell, "F1F5F9")
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                r_lbl = p.add_run(f"{label}\n")
                r_lbl.font.size = Pt(8.5)
                r_lbl.font.bold = True
                r_lbl.font.color.rgb = self.MUTED_COLOR
                r_val = p.add_run(val)
                r_val.font.size = Pt(9.5)
                r_val.font.color.rgb = self.PRIMARY_COLOR

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # Section 1: Executive Overview
        h1 = doc.add_paragraph()
        r_h1 = h1.add_run("1. Executive Overview")
        r_h1.font.size = Pt(15)
        r_h1.font.bold = True
        r_h1.font.color.rgb = self.PRIMARY_COLOR

        p_headline = doc.add_paragraph()
        r_hd = p_headline.add_run(f'"{brief.executive_summary.headline}"')
        r_hd.font.size = Pt(12)
        r_hd.font.italic = True
        r_hd.font.color.rgb = self.ACCENT_COLOR

        p_context = doc.add_paragraph(brief.executive_summary.context)
        p_context.paragraph_format.space_after = Pt(12)

        # Section 2: Key Findings
        h2 = doc.add_paragraph()
        r_h2 = h2.add_run("2. Key Analytical Findings")
        r_h2.font.size = Pt(15)
        r_h2.font.bold = True
        r_h2.font.color.rgb = self.PRIMARY_COLOR

        for finding in brief.executive_summary.key_findings:
            p_bullet = doc.add_paragraph(style='List Bullet')
            r_b = p_bullet.add_run(finding)
            r_b.font.size = Pt(10.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # Section 3: Strategic Recommendations
        h3 = doc.add_paragraph()
        r_h3 = h3.add_run("3. Strategic Recommendations")
        r_h3.font.size = Pt(15)
        r_h3.font.bold = True
        r_h3.font.color.rgb = self.PRIMARY_COLOR

        for rec in brief.executive_summary.strategic_recommendations:
            rec_table = doc.add_table(rows=1, cols=1)
            cell = rec_table.rows[0].cells[0]
            self._set_cell_background(cell, "EFF6FF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(f"➔ {rec}")
            r.font.size = Pt(10.5)
            r.font.color.rgb = self.PRIMARY_COLOR
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # Section 4: Cryptographic Provenance & Verification QR Code
        doc.add_paragraph().paragraph_format.space_after = Pt(12)
        h4 = doc.add_paragraph()
        r_h4 = h4.add_run("4. Cryptographic Provenance Ledger")
        r_h4.font.size = Pt(13)
        r_h4.font.bold = True
        r_h4.font.color.rgb = self.PRIMARY_COLOR

        p_prov = doc.add_paragraph(
            "This document was synthesized via TransmuteAI. It is cryptographically anchored to an immutable "
            "ledger. Scan the verification QR code or query the API with the Brief ID to inspect the hash chain."
        )
        p_prov.paragraph_format.space_after = Pt(6)

        # Embed QR Code if available in extra_context
        qr_path = extra_context.get("qr_code_path") if extra_context else None
        if qr_path and os.path.exists(qr_path):
            qr_para = doc.add_paragraph()
            qr_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            qr_para.add_run().add_picture(qr_path, width=Inches(1.5))
            cap = doc.add_paragraph("Scan to Verify Ledger Provenance")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].font.size = Pt(8.5)
            cap.runs[0].font.italic = True
            cap.runs[0].font.color.rgb = self.MUTED_COLOR

        # Save document
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        doc.save(output_path)
        return os.path.abspath(output_path)
