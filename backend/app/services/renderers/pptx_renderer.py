import os
from typing import Dict, Any, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from app.api.schemas import ContentBriefJSON
from app.services.renderers.base import BaseRenderer

class PPTXRenderer(BaseRenderer):
    """
    Renders ContentBriefJSON into a modern 16:9 presentation deck via python-pptx.
    """

    BG_DARK = RGBColor(15, 23, 42)          # Slate 900
    TEXT_LIGHT = RGBColor(248, 250, 252)    # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)    # Slate 400
    ACCENT_BLUE = RGBColor(59, 130, 246)    # Blue 500
    CARD_BG = RGBColor(30, 41, 59)          # Slate 800

    def render(self, brief: ContentBriefJSON, output_path: str, extra_context: Optional[Dict[str, Any]] = None) -> str:
        prs = Presentation()
        # Set 16:9 widescreen dimensions (13.33 x 7.5 inches)
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        blank_layout = prs.slide_layouts[6]  # Blank slide

        # -------------------------------------------------------------
        # 1. Title Slide
        # -------------------------------------------------------------
        slide_title = prs.slides.add_slide(blank_layout)
        bg_shape = slide_title.shapes.add_shape(1, 0, 0, Inches(13.33), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = self.BG_DARK
        bg_shape.line.color.rgb = self.BG_DARK

        # Main Title Box
        title_box = slide_title.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11.0), Inches(2.2))
        tf = title_box.text_frame
        tf.word_wrap = True
        p_title = tf.paragraphs[0]
        p_title.text = brief.title
        p_title.font.size = Pt(40)
        p_title.font.bold = True
        p_title.font.color.rgb = self.TEXT_LIGHT

        p_sub = tf.add_paragraph()
        p_sub.text = f"Automated Strategic Intelligence | Tone: {brief.tone.capitalize()}"
        p_sub.font.size = Pt(18)
        p_sub.font.color.rgb = self.ACCENT_BLUE
        p_sub.space_before = Pt(14)

        p_meta = tf.add_paragraph()
        p_meta.text = f"Brief ID: {brief.brief_id} • Verified Cryptographic Provenance"
        p_meta.font.size = Pt(12)
        p_meta.font.color.rgb = self.TEXT_MUTED
        p_meta.space_before = Pt(10)

        # -------------------------------------------------------------
        # 2. Executive Summary Slide
        # -------------------------------------------------------------
        slide_exec = prs.slides.add_slide(blank_layout)
        bg2 = slide_exec.shapes.add_shape(1, 0, 0, Inches(13.33), Inches(7.5))
        bg2.fill.solid()
        bg2.fill.fore_color.rgb = self.BG_DARK
        bg2.line.color.rgb = self.BG_DARK

        header_box = slide_exec.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(11.0), Inches(1.0))
        tf_h = header_box.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = "Executive Summary Overview"
        p_h.font.size = Pt(28)
        p_h.font.bold = True
        p_h.font.color.rgb = self.TEXT_LIGHT

        # Left Column: Headline & Summary
        box_left = slide_exec.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(5.2), Inches(4.5))
        tf_l = box_left.text_frame
        tf_l.word_wrap = True
        p_hl = tf_l.paragraphs[0]
        p_hl.text = brief.executive_summary.headline
        p_hl.font.size = Pt(18)
        p_hl.font.bold = True
        p_hl.font.color.rgb = self.ACCENT_BLUE

        p_ctx = tf_l.add_paragraph()
        p_ctx.text = brief.executive_summary.context
        p_ctx.font.size = Pt(13)
        p_ctx.font.color.rgb = self.TEXT_LIGHT
        p_ctx.space_before = Pt(12)

        # Right Column: Key Takeaways Card
        card = slide_exec.shapes.add_shape(1, Inches(6.8), Inches(2.0), Inches(5.3), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = self.CARD_BG
        card.line.color.rgb = self.ACCENT_BLUE

        tf_card = card.text_frame
        tf_card.word_wrap = True
        p_ct = tf_card.paragraphs[0]
        p_ct.text = "Key Strategic Findings"
        p_ct.font.size = Pt(16)
        p_ct.font.bold = True
        p_ct.font.color.rgb = self.ACCENT_BLUE

        for finding in brief.executive_summary.key_findings[:4]:
            p_f = tf_card.add_paragraph()
            p_f.text = f"• {finding}"
            p_f.font.size = Pt(12)
            p_f.font.color.rgb = self.TEXT_LIGHT
            p_f.space_before = Pt(8)

        # -------------------------------------------------------------
        # 3. Content Slides from brief.presentation_outline
        # -------------------------------------------------------------
        for slide_info in brief.presentation_outline:
            slide = prs.slides.add_slide(blank_layout)
            bg = slide.shapes.add_shape(1, 0, 0, Inches(13.33), Inches(7.5))
            bg.fill.solid()
            bg.fill.fore_color.rgb = self.BG_DARK
            bg.line.color.rgb = self.BG_DARK

            # Title
            h_box = slide.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(11.0), Inches(1.0))
            tf_slide = h_box.text_frame
            p_title = tf_slide.paragraphs[0]
            p_title.text = f"{slide_info.slide_number}. {slide_info.title}"
            p_title.font.size = Pt(28)
            p_title.font.bold = True
            p_title.font.color.rgb = self.TEXT_LIGHT

            # Content Card
            content_card = slide.shapes.add_shape(1, Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.5))
            content_card.fill.solid()
            content_card.fill.fore_color.rgb = self.CARD_BG
            content_card.line.color.rgb = self.CARD_BG

            tf_c = content_card.text_frame
            tf_c.word_wrap = True
            tf_c.margin_left = Inches(0.4)
            tf_c.margin_top = Inches(0.4)

            for idx, bullet in enumerate(slide_info.bullet_points):
                p_b = tf_c.paragraphs[0] if idx == 0 else tf_c.add_paragraph()
                p_b.text = f"▸  {bullet}"
                p_b.font.size = Pt(15)
                p_b.font.color.rgb = self.TEXT_LIGHT
                if idx > 0:
                    p_b.space_before = Pt(16)

            # Speaker Notes
            if slide_info.speaker_notes:
                notes_slide = slide.notes_slide
                text_frame = notes_slide.notes_text_frame
                text_frame.text = slide_info.speaker_notes

        # Save Presentation
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        prs.save(output_path)
        return os.path.abspath(output_path)
