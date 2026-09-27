"""
Generates the official 5-Slide Technical Presentation Deck for TransmuteAI.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_DARK = RGBColor(15, 23, 42)        # Slate 900
    TEXT_LIGHT = RGBColor(248, 250, 252)  # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
    ACCENT_BLUE = RGBColor(59, 130, 246)  # Blue 500
    ACCENT_CYAN = RGBColor(6, 182, 212)   # Cyan 500
    ACCENT_EMERALD = RGBColor(16, 185, 129) # Emerald 500
    CARD_BG = RGBColor(30, 41, 59)        # Slate 800

    def add_slide_bg(slide):
        bg = slide.shapes.add_shape(1, 0, 0, Inches(13.33), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.color.rgb = BG_DARK
        return bg

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s1)

    tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.5), Inches(3.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "TransmuteAI"
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CYAN

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Automated Multi-Modal Content Transformation with Cryptographic Provenance"
    p1_sub.font.size = Pt(22)
    p1_sub.font.color.rgb = TEXT_LIGHT
    p1_sub.space_before = Pt(12)

    p1_meta = tf1.add_paragraph()
    p1_meta.text = "Technical Evaluation & Architecture Briefing • Next.js 14 • FastAPI • Google Gemini • Blockchain Ledger"
    p1_meta.font.size = Pt(13)
    p1_meta.font.color.rgb = TEXT_MUTED
    p1_meta.space_before = Pt(20)

    # =========================================================================
    # SLIDE 2: The Core Problem
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s2)

    h2 = s2.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(11.0), Inches(1.0))
    p_h2 = h2.text_frame.paragraphs[0]
    p_h2.text = "The Problem: Semantic Drift & Privacy Leaks"
    p_h2.font.size = Pt(30)
    p_h2.font.bold = True
    p_h2.font.color.rgb = TEXT_LIGHT

    # 3 Cards
    cards_data = [
        ("Multi-Channel Hallucinations", "Generating Docs, Slides, and Scripts independently causes models to output conflicting statistics and divergent recommendations across channels.", RGBColor(239, 68, 68)),
        ("Perimeter PII Exposure", "Raw contracts and memos contain confidential customer emails, SSNs, and credit cards that leak into LLM prompts without perimeter sanitization.", RGBColor(245, 158, 11)),
        ("Absence of Verifiable Trust", "Enterprise deliverables lack cryptographic provenance. Stakeholders cannot prove whether an AI document was altered or authentically generated.", RGBColor(168, 85, 247))
    ]
    for idx, (title, desc, color) in enumerate(cards_data):
        card = s2.shapes.add_shape(1, Inches(1.2 + idx * 3.75), Inches(2.2), Inches(3.45), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = color
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)
        
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(17)
        pt.font.bold = True
        pt.font.color.rgb = color

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(13)
        pd.font.color.rgb = TEXT_LIGHT
        pd.space_before = Pt(14)

    # =========================================================================
    # SLIDE 3: The Breakthrough
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s3)

    h3 = s3.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(11.0), Inches(1.0))
    p_h3 = h3.text_frame.paragraphs[0]
    p_h3.text = "The Solution: Canonical Content Brief JSON"
    p_h3.font.size = Pt(30)
    p_h3.font.bold = True
    p_h3.font.color.rgb = TEXT_LIGHT

    # Left Column: Principles
    left_box = s3.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(5.4), Inches(4.6))
    tf_l = left_box.text_frame
    tf_l.word_wrap = True

    pts = [
        ("Two-Phase Architecture", "Decouples cognitive synthesis from document rendering."),
        ("Deterministic Privacy Perimeter", "Scans and redacts Emails, SSNs, Phones, and Cards before prompting LLM."),
        ("Strict Schema Enforcement", "Pydantic v2 validation guarantees 100% schema integrity."),
        ("Zero Semantic Drift", "All downstream renderers consume the identical JSON representation.")
    ]
    for idx, (label, detail) in enumerate(pts):
        p = tf_l.paragraphs[0] if idx == 0 else tf_l.add_paragraph()
        p.text = f"▸  {label}: {detail}"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_LIGHT
        if idx > 0:
            p.space_before = Pt(16)

    # Right Column: Schema Structure
    schema_card = s3.shapes.add_shape(1, Inches(7.0), Inches(2.0), Inches(5.1), Inches(4.6))
    schema_card.fill.solid()
    schema_card.fill.fore_color.rgb = CARD_BG
    schema_card.line.color.rgb = ACCENT_CYAN
    tf_sc = schema_card.text_frame
    tf_sc.word_wrap = True
    tf_sc.margin_left = Inches(0.4)
    tf_sc.margin_top = Inches(0.4)
    p_sc_title = tf_sc.paragraphs[0]
    p_sc_title.text = "Canonical JSON Schema Architecture"
    p_sc_title.font.size = Pt(16)
    p_sc_title.font.bold = True
    p_sc_title.font.color.rgb = ACCENT_CYAN

    code_snippet = (
        "{\n"
        '  "brief_id": "brief_0123",\n'
        '  "video_package": { storyboard, srt },\n'
        '  "linkedin_post": { hook, bullets, cta },\n'
        '  "social_thread": [ tweet_1, tweet_2 ],\n'
        '  "advisory_document": { risk_matrix },\n'
        '  "infographic": { palette, quadrants },\n'
        '  "executive_summary": { findings },\n'
        '  "presentation_outline": [ slides ]\n'
        "}"
    )
    p_code = tf_sc.add_paragraph()
    p_code.text = code_snippet
    p_code.font.name = "Consolas"
    p_code.font.size = Pt(12)
    p_code.font.color.rgb = RGBColor(147, 197, 253)
    p_code.space_before = Pt(10)

    # =========================================================================
    # SLIDE 4: 7-in-1 Human-Grade Synthesis Engine
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s4)

    h4 = s4.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(11.0), Inches(1.0))
    p_h4 = h4.text_frame.paragraphs[0]
    p_h4.text = "Multi-Format Synthesis: 7 Specialized Deliverables"
    p_h4.font.size = Pt(30)
    p_h4.font.bold = True
    p_h4.font.color.rgb = TEXT_LIGHT

    deliverables = [
        ("1. Video Package", "Full screenplay, scene cinematography directions, audio cues, b-roll, and .SRT broadcast subtitles."),
        ("2. LinkedIn Post", "Mobile-optimized thought leadership post with curiosity hooks, white space, bullet takeaways, and CTA."),
        ("3. Twitter/X Thread", "High-velocity numbered tweet sequence (1/N) structured for viral distribution and retention."),
        ("4. Strategic Advisory", "Formal McKinsey-style advisory memo featuring quantitative Risk Matrices and 30-60-90 day Roadmaps."),
        ("5. Infographic Spec", "Layout blueprint, color palette guidelines, hero stat quadrants, and chart visual directives."),
        ("6. Executive Summary", "Corporate executive briefing Word DOCX with structured tables, callouts, and embedded QR code."),
        ("7. Presentation Deck", "Widescreen 16:9 PowerPoint slide deck with dark executive theme and comprehensive speaker notes.")
    ]

    for idx, (name, desc) in enumerate(deliverables):
        col = 0 if idx < 4 else 1
        row_idx = idx if col == 0 else idx - 4
        
        box = s4.shapes.add_shape(1, Inches(1.2 + col * 5.6), Inches(2.0 + row_idx * 1.2), Inches(5.3), Inches(1.05))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ACCENT_BLUE if col == 0 else ACCENT_CYAN
        tf_b = box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = Inches(0.25)
        tf_b.margin_top = Inches(0.15)

        pb1 = tf_b.paragraphs[0]
        pb1.text = name
        pb1.font.size = Pt(13)
        pb1.font.bold = True
        pb1.font.color.rgb = ACCENT_CYAN

        pb2 = tf_b.add_paragraph()
        pb2.text = desc
        pb2.font.size = Pt(10)
        pb2.font.color.rgb = TEXT_LIGHT
        pb2.space_before = Pt(2)

    # =========================================================================
    # SLIDE 5: Cryptographic Provenance & Infrastructure
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s5)

    h5 = s5.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(11.0), Inches(1.0))
    p_h5 = h5.text_frame.paragraphs[0]
    p_h5.text = "Cryptographic Trust, Scalability & Free Deployment"
    p_h5.font.size = Pt(30)
    p_h5.font.bold = True
    p_h5.font.color.rgb = TEXT_LIGHT

    # 3 Summary Cards
    s5_cards = [
        ("Blockchain Provenance", [
            "SHA-256 fingerprinting of all inputs and outputs",
            "Chained block ledger with previous_hash validation",
            "Tamper-evident verification algorithms",
            "Embedded dynamic verification QR codes"
        ], ACCENT_EMERALD),
        ("Scalable Execution", [
            "Celery + Redis distributed task queue for enterprise scale",
            "Direct BackgroundTasks mode for serverless/lightweight hosting",
            "SlowAPI client IP rate limiting",
            "Async connection pooling with SQLAlchemy"
        ], ACCENT_BLUE),
        ("100% Free Cloud Hosting", [
            "Frontend: Vercel (Hobby Tier - $0)",
            "Backend: Render.com / Koyeb Web Service ($0)",
            "LLM: Google Gemini Free Tier via Google AI Studio",
            "Database: Neon.tech Serverless Postgres ($0)"
        ], ACCENT_CYAN)
    ]

    for idx, (title, points, color) in enumerate(s5_cards):
        card = s5.shapes.add_shape(1, Inches(1.2 + idx * 3.75), Inches(2.2), Inches(3.45), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = color
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)

        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(17)
        pt.font.bold = True
        pt.font.color.rgb = color

        for p_idx, point in enumerate(points):
            pp = tf.add_paragraph()
            pp.text = f"• {point}"
            pp.font.size = Pt(12)
            pp.font.color.rgb = TEXT_LIGHT
            pp.space_before = Pt(10)

    output_file = "transmuteai_technical_presentation.pptx"
    prs.save(output_file)
    print(f"Technical Presentation created: {output_file}")

if __name__ == "__main__":
    create_presentation()
